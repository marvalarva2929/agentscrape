"""Extract client-supplied "priority" URLs from an uploaded spreadsheet.

A client's spreadsheet is unstructured: URLs can sit in any cell of any
worksheet, mixed with program names, contact info, or anything else. This
module never assumes a column name, a column position, or a single
worksheet - it scans every populated cell and keeps only what parses as an
http(s) URL. No LLM involved: URL recognition is a deterministic regex plus
existing canonicalization, not a judgment call.
"""

from __future__ import annotations

import csv
import io
import logging
import re

from openpyxl import load_workbook

from ..urls import canonicalize

log = logging.getLogger("agentscrape.discovery.priority_urls")

# A generous http(s) URL matcher: stops at whitespace or a small set of
# characters that are never part of a URL but commonly sit right after one in
# free text (closing punctuation, quotes).
_URL_RE = re.compile(r"https?://[^\s<>\"')\]]+", re.IGNORECASE)


def _urls_in_text(value: object) -> list[str]:
    if not isinstance(value, str) or "http" not in value.lower():
        return []
    return [match.group(0).rstrip(".,;:!?") for match in _URL_RE.finditer(value)]


def _normalize_and_dedupe(raw_urls: list[str]) -> list[str]:
    """Canonicalize with the crawler's own URL logic and drop duplicates,
    keeping first-seen order so the client sees a stable, predictable list."""
    seen: dict[str, None] = {}
    for raw in raw_urls:
        canonical = canonicalize(raw) or raw
        seen.setdefault(canonical, None)
    return list(seen)


def extract_urls_from_csv(content: bytes) -> list[str]:
    """Every recognizable URL in every field of every row. Column-agnostic."""
    text = content.decode("utf-8-sig", errors="replace")
    if not text.strip():
        return []
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    found: list[str] = []
    for row in csv.reader(io.StringIO(text), dialect):
        for field in row:
            found.extend(_urls_in_text(field))
    return _normalize_and_dedupe(found)


def extract_urls_from_xlsx(content: bytes) -> list[str]:
    """Every recognizable URL in every populated cell across every worksheet,
    including a hyperlink's real target when the cell only displays friendly
    text (e.g. "Internal Medicine" linking to a roster page)."""
    found: list[str] = []
    # Not read_only: openpyxl's read-only cells drop per-cell hyperlink access
    # entirely, which is the only way to recover a link hidden behind display
    # text that isn't itself a URL (confirmed against openpyxl's own behavior).
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        found.extend(_urls_in_text(cell.value))
                    link = getattr(cell, "hyperlink", None)
                    target = getattr(link, "target", None)
                    if target:
                        found.extend(_urls_in_text(target))
    finally:
        workbook.close()
    return _normalize_and_dedupe(found)


def google_sheet_csv_export_url(sheet_url: str) -> str | None:
    """The CSV-export URL for a Google Sheet shared as "anyone with the link
    can view", or None if `sheet_url` isn't recognizable as a Sheets URL.

    v1 scope: one worksheet/tab per link (whichever `gid` the link encodes,
    or the first tab if none) - no Sheets API, no OAuth, no multi-tab
    enumeration. A private, unshared sheet cannot be read this way.
    """
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", sheet_url)
    if not match:
        return None
    sheet_id = match.group(1)
    gid_match = re.search(r"[?&#]gid=(\d+)", sheet_url)
    gid = gid_match.group(1) if gid_match else "0"
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
